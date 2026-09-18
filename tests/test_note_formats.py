# -*- coding: utf-8 -*-
"""笔记双格式（markdown / richtext）：数据模型、迁移、编辑器、新建选格式、页面切换。"""
import os
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from sqlalchemy import text

import zhixing.model.infrastructure.db as dbmod


def _mk_app():
    from PySide6.QtWidgets import QApplication
    from qfluent_core import ThemeManager as ThemeEngine
    app = QApplication.instance() or QApplication([])
    inst = getattr(ThemeEngine, "instance", None)
    if not isinstance(inst, ThemeEngine):
        eng = ThemeEngine()
        eng.apply("青竹", "light", "#0D9488")
        ThemeEngine.install(eng)   # 装成单例（原为属性赋值）
    return app


def _make_note_service(db):
    from zhixing.core.event_bus import EventBus
    from zhixing.model.application.note_service import NoteService
    from zhixing.model.infrastructure.fts import FTSService
    from zhixing.model.infrastructure.repositories import (
        NoteLinkRepository, NoteRepository, TagRepository,
    )
    fts = FTSService(db)
    return NoteService(db, EventBus(), NoteRepository(db, fts),
                       NoteLinkRepository(db), TagRepository(db))


class _FakeNoteService:
    """NotePage 切换测试用内存服务。"""

    def __init__(self, notes):
        self.notes = {n.id: n for n in notes}

    def get(self, note_id):
        return self.notes.get(note_id)

    def list(self, **kw):
        return list(self.notes.values())

    def folder_titles(self):
        return []

    def backlinks(self, note_id):
        return []

    def resolve(self, title):
        return None

    def recent(self, limit=5):
        return list(self.notes.values())[:limit]

    def save(self, note_id, **kw):
        n = self.notes.get(note_id)
        if n is None:
            return None
        if kw.get("title") is not None:
            n.title = kw["title"]
        if kw.get("content_md") is not None:
            n.content_md = kw["content_md"]
        return n


class TestNoteFormatModel(unittest.TestCase):
    def test_note_entity_defaults_to_markdown(self):
        from zhixing.model.domain.entities import Note
        self.assertEqual(Note().format, "markdown")
        self.assertEqual(Note(format="richtext").format, "richtext")

    def test_note_row_has_format_column(self):
        from zhixing.model.infrastructure.models import NoteRow, SCHEMA_VERSION
        col = NoteRow.__table__.c.format
        self.assertFalse(col.nullable)
        self.assertEqual(SCHEMA_VERSION, 12)  # v12: note_link.dst_title 建索引（重命名加速）


class TestNoteFormatLabels(unittest.TestCase):
    def test_format_label_mapping(self):
        from zhixing.model.application.note_service import (
            NOTE_FORMAT_LABELS, format_from_label, format_label,
        )
        self.assertEqual([f for f, _ in NOTE_FORMAT_LABELS],
                         ["markdown", "richtext", "word", "excel", "link"])
        self.assertEqual(format_from_label("富文本"), "richtext")
        self.assertEqual(format_from_label("Markdown"), "markdown")
        self.assertEqual(format_from_label("Word"), "word")
        self.assertEqual(format_from_label("Excel"), "excel")
        self.assertEqual(format_from_label("链接"), "link")
        self.assertEqual(format_from_label(None), "markdown")
        self.assertEqual(format_label("richtext"), "富文本")
        self.assertEqual(format_label("markdown"), "Markdown")
        self.assertEqual(format_label("word"), "Word")
        self.assertEqual(format_label("excel"), "Excel")
        self.assertEqual(format_label("link"), "链接")


class TestV3Migration(unittest.TestCase):
    def setUp(self):
        self._home = tempfile.mkdtemp()
        os.environ["ZHIXING_HOME"] = self._home

    def test_v3_adds_format_column_and_bumps_version(self):
        """V2 老库（note 无 format 列）应自动加列并升到最新版，且幂等。"""
        dbpath = str(dbmod.db_path())
        conn = sqlite3.connect(dbpath)
        try:
            conn.execute("CREATE TABLE settings (key TEXT PRIMARY KEY, value TEXT DEFAULT '')")
            conn.execute("INSERT INTO settings(key, value) VALUES('schema_version', '2')")
            conn.execute(
                "CREATE TABLE note (id INTEGER PRIMARY KEY AUTOINCREMENT, "
                "folder_id INTEGER, title TEXT NOT NULL DEFAULT '无标题', "
                "content_md TEXT DEFAULT '', pinned BOOLEAN DEFAULT 0, "
                "word_count INTEGER DEFAULT 0, deleted_at DATETIME, created_at DATETIME, "
                "updated_at DATETIME)")
            conn.commit()
        finally:
            conn.close()

        db = dbmod.Database()
        try:
            s = db.session()
            try:
                cols = [r[1] for r in s.execute(text("PRAGMA table_info(note)"))]
                self.assertIn("format", cols)
                v = s.execute(text(
                    "SELECT value FROM settings WHERE key='schema_version'")).scalar()
                self.assertEqual(int(v), dbmod.SCHEMA_VERSION)
            finally:
                s.close()
            self.assertEqual(db.migrate(), [], "迁移应幂等")
        finally:
            db.engine.dispose()


class TestNoteServiceFormat(unittest.TestCase):
    def setUp(self):
        os.environ["ZHIXING_HOME"] = tempfile.mkdtemp()
        self.db = dbmod.Database()
        self.svc = _make_note_service(self.db)

    def tearDown(self):
        try:
            self.db.engine.dispose()
        except Exception:
            pass

    def test_create_with_format_roundtrips(self):
        n = self.svc.create(title="富文本笔记", content_md="<p>你好</p>", format="richtext")
        self.assertEqual(n.format, "richtext")
        fresh = self.svc.get(n.id)
        self.assertEqual(fresh.format, "richtext")
        self.assertEqual(fresh.content_md, "<p>你好</p>")

    def test_create_defaults_to_markdown_and_sanitizes_bad_format(self):
        m = self.svc.create(title="Markdown 笔记", content_md="# 标题")
        self.assertEqual(m.format, "markdown")
        bad = self.svc.create(title="非法格式回落", format="weird")
        self.assertEqual(bad.format, "markdown")


class TestMarkdownToolbar(unittest.TestCase):
    def setUp(self):
        _mk_app()
        from zhixing.view.components.markdown_editor import MarkdownEditor
        self.editor = MarkdownEditor()

    def tearDown(self):
        self.editor.deleteLater()

    def test_toolbar_inserts_markdown_syntax(self):
        self.editor.set_content("")
        self.editor.insert_bold()
        self.assertEqual(self.editor.content(), "**加粗**")

        self.editor.set_content("")
        self.editor.insert_italic()
        self.assertEqual(self.editor.content(), "*斜体*")

        self.editor.set_content("")
        self.editor.insert_heading(2)
        self.assertEqual(self.editor.content(), "## ")

        self.editor.set_content("")
        self.editor.insert_unordered_list()
        self.assertEqual(self.editor.content(), "- ")

        self.editor.set_content("")
        self.editor.insert_ordered_list()
        self.assertEqual(self.editor.content(), "1. ")

        self.editor.set_content("")
        self.editor.insert_link()
        self.assertEqual(self.editor.content(), "[文字](url)")

        self.editor.set_content("")
        self.editor.insert_code_block()
        self.assertIn("```", self.editor.content())
        self.assertIn("代码", self.editor.content())

    def test_toolbar_wraps_selection(self):
        from PySide6.QtGui import QTextCursor
        self.editor.set_content("你好")
        cur = self.editor.editor.textCursor()
        cur.setPosition(0)
        cur.setPosition(2, QTextCursor.KeepAnchor)
        self.editor.editor.setTextCursor(cur)
        self.editor.insert_bold()
        self.assertEqual(self.editor.content(), "**你好**")


class TestRichTextEditor(unittest.TestCase):
    def setUp(self):
        _mk_app()
        from zhixing.view.components.richtext_editor import RichTextEditor
        self.editor = RichTextEditor()

    def tearDown(self):
        self.editor.deleteLater()

    def test_set_html_and_to_html_roundtrip(self):
        self.editor.set_html("<p>你好 <b>加粗</b></p>")
        html = self.editor.to_html()
        self.assertIn("加粗", html)
        self.assertIn("font-weight:700", html, "粗体应序列化为 font-weight:700")
        # to_html 结果可再 set_html 回填，不崩
        self.editor.set_html(html)
        self.assertIn("加粗", self.editor.to_html())

    def test_toolbar_bold_produces_bold_tag(self):
        from PySide6.QtGui import QTextCursor
        self.editor.edit.setPlainText("加粗文本")
        cur = self.editor.edit.textCursor()
        cur.select(QTextCursor.Document)
        self.editor.edit.setTextCursor(cur)
        self.editor.toggle_bold()
        self.assertIn("font-weight:700", self.editor.to_html())

    def test_toolbar_heading_produces_heading_tag(self):
        self.editor.edit.setPlainText("标题行")
        self.editor.set_heading(2)
        self.assertIn("<h2", self.editor.to_html())


class TestNotePageEditorSwitch(unittest.TestCase):
    def test_select_note_switches_editor_by_format(self):
        _mk_app()
        from zhixing.model.domain.entities import Note
        from zhixing.view.pages.note_page import NotePage

        md = Note(id=1, title="MD", content_md="# 标题\n正文")
        rt = Note(id=2, title="RT", content_md="<p>你好 <b>加粗</b></p>", format="richtext")
        page = NotePage(_FakeNoteService([md, rt]), None)
        try:
            page.select_note(2)
            self.assertIs(page.editor.currentWidget(), page.rt_editor)
            self.assertIn("加粗", page.rt_editor.to_html())

            page.select_note(1)
            self.assertIs(page.editor.currentWidget(), page.md_editor)
            self.assertEqual(page.md_editor.content(), "# 标题\n正文")
        finally:
            page.deleteLater()


class _TitleFlowNoteService:
    """需求④：新建/重命名/自动保存标题流冒烟用的内存 note_service。"""

    def __init__(self):
        self.notes = {}
        self._next_id = 1

    def get(self, note_id):
        return self.notes.get(note_id)

    def list(self, **kw):
        return list(self.notes.values())

    def folder_titles(self):
        return []

    def ensure_default_folder(self):
        return None

    def backlinks(self, note_id):
        return []

    def resolve(self, title):
        return None

    def recent(self, limit=5):
        return list(self.notes.values())[:limit]

    def save(self, note_id, **kw):
        n = self.notes.get(note_id)
        if n is None:
            return None
        if kw.get("title") is not None:
            n.title = kw["title"]
        if kw.get("content_md") is not None:
            n.content_md = kw["content_md"]
        return n

    def create(self, title="", content_md="", folder_id=None, **kw):
        from zhixing.model.domain.entities import Note
        n = Note(id=self._next_id, title=title, content_md=content_md, folder_id=folder_id,
                 format=kw.get("format", "markdown"))
        self._next_id += 1
        self.notes[n.id] = n
        return n


class TestNoteTitleFromUserInput(unittest.TestCase):
    """需求④：新建/重命名以用户输入为标题，自动保存不再按内容首行改写标题。"""

    def _mk_page(self):
        _mk_app()
        from zhixing.view.pages.note_page import NotePage
        svc = _TitleFlowNoteService()
        return NotePage(svc, None), svc

    def test_markdown_editor_emit_save_no_first_line_title(self):
        _mk_app()
        from zhixing.view.components.markdown_editor import MarkdownEditor
        ed = MarkdownEditor()
        try:
            captured = []
            ed.saveRequested.connect(lambda t, c: captured.append((t, c)))
            ed.editor.setPlainText("# 我的标题\n正文")   # 触发 textChanged → _dirty=True
            ed.commit()   # 触发 _emit_save
            self.assertTrue(captured)
            self.assertEqual(captured[-1][0], "", "编辑器不应再按首行推导标题")
            self.assertEqual(captured[-1][1], "# 我的标题\n正文")
        finally:
            ed.deleteLater()

    def test_auto_save_preserves_user_title(self):
        from zhixing.model.domain.entities import Note
        page, svc = self._mk_page()
        svc.notes[1] = Note(id=1, title="用户命名", content_md="# 内容首行\n正文")
        try:
            page.select_note(1)
            page._auto_save("忽略的首行", "# 内容首行\n新正文")
            self.assertEqual(svc.notes[1].title, "用户命名",
                             "自动保存不应覆盖用户输入的标题")
            self.assertEqual(svc.notes[1].content_md, "# 内容首行\n新正文")
        finally:
            page.deleteLater()

    def test_prompt_new_note_uses_input_title(self):
        from unittest import mock
        page, svc = self._mk_page()
        try:
            with mock.patch("zhixing.view.pages.note_page.NoteCreateDialog.get_note",
                            return_value=("我的新笔记", "markdown", "")):
                page._prompt_new_note(None)
            created = list(svc.notes.values())
            self.assertEqual(len(created), 1)
            self.assertEqual(created[0].title, "我的新笔记")
        finally:
            page.deleteLater()

    def test_prompt_new_note_empty_name_falls_back(self):
        from unittest import mock
        page, svc = self._mk_page()
        try:
            with mock.patch("zhixing.view.pages.note_page.NoteCreateDialog.get_note",
                            return_value=("   ", "markdown", "")):
                page._prompt_new_note(None)
            created = list(svc.notes.values())
            self.assertEqual(created[0].title, "未命名笔记",
                             "空名称应回退到明确默认标题")
        finally:
            page.deleteLater()

    def test_rename_note_uses_input_title_and_syncs(self):
        from unittest import mock
        from zhixing.model.domain.entities import Note
        page, svc = self._mk_page()
        svc.notes[1] = Note(id=1, title="旧标题", content_md="正文")
        try:
            page.reload_folders()
            page.select_note(1)
            with mock.patch("zhixing.view.pages.note_page.UInputDialog.get_text",
                            return_value=("新标题", True)):
                page._rename_note(1)
            self.assertEqual(svc.notes[1].title, "新标题",
                             "重命名应以用户输入为最终标题")
        finally:
            page.deleteLater()


class TestNoteServiceNewFormats(unittest.TestCase):
    """需求⑦：Word/Excel/链接 三类的 format 持久化与内容承载。"""

    def setUp(self):
        os.environ["ZHIXING_HOME"] = tempfile.mkdtemp()
        self.db = dbmod.Database()
        self.svc = _make_note_service(self.db)

    def tearDown(self):
        try:
            self.db.engine.dispose()
        except Exception:
            pass

    def test_word_excel_link_roundtrip(self):
        cases = [
            ("word", "/tmp/a.docx"),
            ("excel", "/tmp/b.xlsx"),
            ("link", "https://example.com"),
        ]
        for fmt, content in cases:
            n = self.svc.create(title=fmt, content_md=content, format=fmt)
            self.assertEqual(n.format, fmt)
            fresh = self.svc.get(n.id)
            self.assertEqual(fresh.format, fmt, "重开/重读后 format 应正确恢复")
            self.assertEqual(fresh.content_md, content)

    def test_unknown_format_still_falls_back(self):
        n = self.svc.create(title="未知", format="pdf")
        self.assertEqual(n.format, "markdown")


class TestV6Migration(unittest.TestCase):
    def setUp(self):
        self._home = tempfile.mkdtemp()
        os.environ["ZHIXING_HOME"] = self._home

    def test_v6_normalizes_empty_format(self):
        """V5 老库 note/note_revision 的空 format 应归一为 markdown。"""
        dbpath = str(dbmod.db_path())
        conn = sqlite3.connect(dbpath)
        try:
            conn.execute("CREATE TABLE settings (key TEXT PRIMARY KEY, value TEXT DEFAULT '')")
            conn.execute("INSERT INTO settings(key, value) VALUES('schema_version', '5')")
            conn.execute(
                "CREATE TABLE note (id INTEGER PRIMARY KEY AUTOINCREMENT, folder_id INTEGER, "
                "title TEXT NOT NULL DEFAULT '无标题', content_md TEXT DEFAULT '', "
                "format TEXT NOT NULL DEFAULT 'markdown', pinned BOOLEAN DEFAULT 0, "
                "word_count INTEGER DEFAULT 0, deleted_at DATETIME, created_at DATETIME, "
                "updated_at DATETIME)")
            conn.execute(
                "CREATE TABLE note_revision (id INTEGER PRIMARY KEY AUTOINCREMENT, "
                "note_id INTEGER NOT NULL, title VARCHAR DEFAULT '', content_md TEXT DEFAULT '', "
                "format VARCHAR DEFAULT 'markdown', created_at DATETIME)")
            conn.execute("INSERT INTO note(id, title, content_md, format) VALUES(1, '空格式', 'x', '')")
            conn.execute("INSERT INTO note_revision(note_id, title, content_md, format) VALUES(1, 't', 'c', NULL)")
            conn.commit()
        finally:
            conn.close()

        db = dbmod.Database()
        try:
            s = db.session()
            try:
                fmt = s.execute(text("SELECT format FROM note WHERE id=1")).scalar()
                self.assertEqual(fmt, "markdown")
                rfmt = s.execute(text("SELECT format FROM note_revision WHERE note_id=1")).scalar()
                self.assertEqual(rfmt, "markdown")
            finally:
                s.close()
        finally:
            db.engine.dispose()


class TestNotePreviews(unittest.TestCase):
    """需求⑦：Word/Excel/链接 的解析与只读预览。"""

    def setUp(self):
        self._tmp = tempfile.mkdtemp()
        self.addCleanup(lambda: None)

    def test_parse_docx_and_xlsx(self):
        from zhixing.view.components.note_previews import parse_docx_to_html, parse_xlsx_to_html

        docx_path = Path(self._tmp) / "a.docx"
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello 世界")
        doc.save(str(docx_path))
        self.assertIn("Hello 世界", parse_docx_to_html(str(docx_path)))

        xlsx_path = Path(self._tmp) / "b.xlsx"
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "名称"
        ws["B1"] = "数量"
        ws["A2"] = "苹果"
        wb.save(str(xlsx_path))
        html = parse_xlsx_to_html(str(xlsx_path))
        self.assertIn("名称", html)
        self.assertIn("苹果", html)

    def test_classify_and_text_preview(self):
        from zhixing.view.components.note_previews import (
            classify_target, preview_text_file,
        )
        self.assertEqual(classify_target("https://example.com"), ("web", "https://example.com"))
        self.assertEqual(classify_target("www.example.com"), ("web", "www.example.com"))
        self.assertEqual(classify_target("/tmp/file.txt"), ("local", "/tmp/file.txt"))
        p = Path(self._tmp) / "note.txt"
        p.write_text("本地预览内容", encoding="utf-8")
        self.assertIn("本地预览内容", preview_text_file(str(p)))

    def test_preview_views_render_readonly(self):
        _mk_app()
        from zhixing.view.components.note_previews import (
            ExcelPreviewView, LinkPreviewView, WordPreviewView,
        )
        from docx import Document
        from openpyxl import Workbook

        docx_path = Path(self._tmp) / "v.docx"
        doc = Document()
        doc.add_paragraph("只读 Word")
        doc.save(str(docx_path))
        w = WordPreviewView()
        w.set_path(str(docx_path))
        self.assertIn("只读 Word", w.browser.toPlainText())

        xlsx_path = Path(self._tmp) / "v.xlsx"
        wb = Workbook()
        wb.active["A1"] = "只读 Excel"
        wb.save(str(xlsx_path))
        e = ExcelPreviewView()
        e.set_path(str(xlsx_path))
        self.assertIn("只读 Excel", e.browser.toPlainText())

        txt = Path(self._tmp) / "link.txt"
        txt.write_text("链接预览", encoding="utf-8")
        lk = LinkPreviewView()
        lk.set_target(str(txt))
        # 链接详情预览区已移除：列表即全部信息，目标存于链接列表
        self.assertEqual([x["target"] for x in lk.links()], [str(txt)])
        lk.set_target("https://example.com")
        # 改造为多链接后，目标展示在链接列表里（target_label 已移除）
        targets = [x["target"] for x in lk.links()]
        self.assertIn("example.com", " ".join(targets))

        w.deleteLater(); e.deleteLater(); lk.deleteLater()


class TestCtrlShiftNNewNotePath(unittest.TestCase):
    """需求④返工：Ctrl+Shift+N（及控制器入口）统一走 NoteCreateDialog（名称+五类型）。"""

    def test_controller_new_note_uses_note_create_dialog(self):
        from types import SimpleNamespace
        from unittest import mock
        from zhixing.controller.app_controller import AppController

        created = []

        class _Svc:
            def create(self, title="", content_md="", folder_id=None,
                       format="markdown", **kw):
                from zhixing.model.domain.entities import Note
                n = Note(id=1, title=title, content_md=content_md,
                         folder_id=folder_id, format=format)
                created.append(n)
                return n

        fake = SimpleNamespace(
            ctx=SimpleNamespace(note_service=_Svc()),
            main=None,
        )
        fake._refresh_notes = lambda: None
        fake._show_note = lambda nid: None

        with mock.patch("zhixing.controller.app_controller.NoteCreateDialog.get_note",
                        return_value=("会议记录", "word", "/tmp/meeting.docx")):
            AppController._prompt_create_note(fake)

        self.assertEqual(len(created), 1)
        self.assertEqual(created[0].title, "会议记录")
        self.assertEqual(created[0].format, "word")
        self.assertEqual(created[0].content_md, "/tmp/meeting.docx")

    def test_ctrl_shift_n_shortcut_wired_to_new_note_no_legacy(self):
        import importlib
        m = importlib.import_module("zhixing.controller.app_controller")
        with open(m.__file__, encoding="utf-8") as f:
            src = f.read()
        # Ctrl+Shift+N → _new_note → _prompt_create_note（NoteCreateDialog）
        self.assertIn('QKeySequence("Ctrl+Shift+N")', src)
        self.assertIn("_new_note_shortcut.activated.connect(self._new_note)", src)
        self.assertIn("NoteCreateDialog.get_note", src)
        self.assertIn("def _new_note(self):\n        self._prompt_create_note()", src)
        # 旧路径（格式选择 + 硬编码无标题）与死接线已清理
        self.assertNotIn("_choose_note_format", src)
        self.assertNotIn('title="无标题"', src)
        self.assertNotIn("createNoteRequested.connect(self._new_note)", src)
        self.assertNotIn("addNoteRequested.connect(self._new_note_in_folder)", src)


if __name__ == "__main__":
    unittest.main(verbosity=2)
