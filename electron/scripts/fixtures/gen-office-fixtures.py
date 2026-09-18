"""生成 officecheck.mjs 需要的 docx / xlsx 夹具。

用法：python gen-office-fixtures.py <docx 路径> <xlsx 路径>

放在仓库里而不是 /tmp：这个脚本原先写死在 /tmp/gen-office-fixtures.py，
系统一清理临时目录，整条 Office 预览验证就再也跑不起来（表现为 python
报 "can't open file"）。夹具生成必须跟验证脚本一起进版本库。
"""
import sys

from docx import Document
from openpyxl import Workbook


def make_docx(path: str) -> None:
    doc = Document()
    doc.add_paragraph("知行 Office 预览验证段落")
    doc.add_paragraph("第二段落内容")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "单元格A"
    table.rows[0].cells[1].text = "单元格B"
    doc.save(path)


def make_xlsx(path: str) -> None:
    wb = Workbook()
    first = wb.active
    first.title = "验证表"
    first.append(["名称", "数量"])
    first.append(["知行", 100])
    second = wb.create_sheet("第二表")
    second.append(["备注", "内容"])
    second.append(["第二表数据", 200])
    wb.save(path)


if __name__ == "__main__":
    make_docx(sys.argv[1])
    make_xlsx(sys.argv[2])
